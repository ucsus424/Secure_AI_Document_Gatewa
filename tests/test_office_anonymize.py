import importlib.util
import io
import re
import sqlite3
import tempfile
import unittest
from pathlib import Path

import app
from core.database import load_mapping
from core.detector import DetectorError
from core.models import Entity


DOCX_AVAILABLE = importlib.util.find_spec("docx") is not None
OPENPYXL_AVAILABLE = importlib.util.find_spec("openpyxl") is not None
PPTX_AVAILABLE = importlib.util.find_spec("pptx") is not None


class OfficeFakeDetector:
    def detect_text(self, text):
        return [
            Entity(type="PERSON", value="Петров А.В.", confidence=0.99),
            Entity(type="COMPANY", value='ООО "Альфа Телеком"', confidence=0.99),
            Entity(type="OBJECT", value="ВРН-Север-17", confidence=0.99),
            Entity(type="IP", value="10.74.16.1/24", confidence=0.99),
        ]


class FailingDetector:
    def detect_text(self, text):
        raise DetectorError("invalid model response")


class PunctuationVariantDetector:
    def detect_text(self, text):
        return [Entity(type="PERSON", value="ИП Некрылова Г.В.", confidence=0.99)]


class PersonOnlyDetector:
    def detect_text(self, text):
        return [Entity(type="PERSON", value="Петров А.В.", confidence=0.99)]


def run_anonymize(input_path, detector):
    """Run `anonymize` with a temporary database next to the input file."""
    db_path = input_path.parent / "test.db"
    stdout = io.StringIO()
    stderr = io.StringIO()
    code = app.run(
        ["anonymize", "--db", str(db_path), str(input_path)],
        detector_factory=lambda args: detector,
        stdout=stdout,
        stderr=stderr,
    )
    match = re.search(r"Document ID: ([0-9A-F]{8})", stdout.getvalue())
    document_id = match.group(1) if match else None
    return code, stdout.getvalue(), stderr.getvalue(), document_id, db_path


def safe_path_for(input_path, document_id):
    return input_path.with_name(f"{input_path.stem}_SAFE_{document_id}{input_path.suffix}")


def count_documents(db_path):
    if not db_path.exists():
        return 0
    with sqlite3.connect(db_path) as connection:
        return connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0]


@unittest.skipUnless(DOCX_AVAILABLE, "python-docx is unavailable")
class DocxAnonymizeTests(unittest.TestCase):
    def test_anonymize_writes_safe_docx_and_saves_values_to_database(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            input_path = Path(tmp) / "report.docx"
            document = Document()
            document.add_paragraph("Инженер Петров А.В. прибыл на узел ВРН-Север-17.")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = 'Заказчик ООО "Альфа Телеком"'
            table.cell(0, 1).text = "IP 10.74.16.1/24"
            document.save(input_path)

            code, stdout, stderr, document_id, db_path = run_anonymize(input_path, OfficeFakeDetector())

            self.assertEqual(code, 0)
            self.assertEqual(stderr, "")
            safe_path = safe_path_for(input_path, document_id)
            self.assertTrue(safe_path.exists())
            self.assertEqual(list(Path(tmp).glob("*.mapping.json")), [])

            safe_document = Document(safe_path)
            safe_text = "\n".join(_docx_text(safe_document))
            self.assertNotIn("Петров А.В.", safe_text)
            self.assertNotIn('ООО "Альфа Телеком"', safe_text)
            self.assertNotIn("ВРН-Север-17", safe_text)
            self.assertNotIn("10.74.16.1/24", safe_text)
            self.assertIn("⟦PERSON_001⟧", safe_text)
            self.assertIn("⟦COMPANY_001⟧", safe_text)
            self.assertIn("⟦OBJECT_001⟧", safe_text)
            self.assertIn("⟦IP_001⟧", safe_text)

            mapping = load_mapping(document_id, db_path)
            self.assertEqual(mapping.source_name, "report.docx")
            self.assertEqual(mapping.entities["PERSON_001"], "Петров А.В.")

    def test_docx_anonymize_replaces_punctuation_variant(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            input_path = Path(tmp) / "report.docx"
            document = Document()
            document.add_paragraph("на объекте ИП Некрылова Г.В,, расположенного.")
            document.add_paragraph("Проектное решение согласовать с ИП Некрылова Г.В.")
            document.save(input_path)

            code, stdout, stderr, document_id, db_path = run_anonymize(input_path, PunctuationVariantDetector())

            self.assertEqual(code, 0)
            safe_document = Document(safe_path_for(input_path, document_id))
            safe_text = "\n".join(_docx_text(safe_document))
            self.assertEqual(safe_text.count("⟦PERSON_001⟧"), 2)
            self.assertNotIn("ИП Некрылова", safe_text)
            self.assertEqual(stderr, "")

    def test_docx_masks_multiline_labeled_address_missed_by_model(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            input_path = Path(tmp) / "report.docx"
            document = Document()
            document.add_paragraph("Ответственный: Петров А.В.")
            document.add_paragraph("Юридический адрес: 123456, Московская область,")
            document.add_paragraph("г. Примерск, ул. Тестовая, д. 1,")
            document.add_paragraph("помещение 2, офис 3")
            document.add_paragraph("e-mail: user@example.com")
            document.save(input_path)

            code, stdout, stderr, document_id, db_path = run_anonymize(input_path, PersonOnlyDetector())

            self.assertEqual(code, 0)
            safe_document = Document(safe_path_for(input_path, document_id))
            safe_text = "\n".join(_docx_text(safe_document))
            self.assertNotIn("123456", safe_text)
            self.assertNotIn("Примерск", safe_text)
            self.assertNotIn("Тестовая", safe_text)
            self.assertNotIn("user@example.com", safe_text)
            self.assertIn("⟦ADDRESS_", safe_text)
            self.assertIn("⟦EMAIL_", safe_text)
            self.assertEqual(stderr, "")

    def test_docx_detector_failure_writes_nothing(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            input_path = Path(tmp) / "report.docx"
            document = Document()
            document.add_paragraph("Пароль Qwerty-2026!")
            document.save(input_path)

            code, stdout, stderr, document_id, db_path = run_anonymize(input_path, FailingDetector())

            self.assertNotEqual(code, 0)
            self.assertEqual(stdout, "")
            self.assertEqual(list(Path(tmp).glob("report_SAFE_*")), [])
            self.assertEqual(count_documents(db_path), 0)
            self.assertNotIn("Qwerty-2026!", stderr)


@unittest.skipUnless(OPENPYXL_AVAILABLE, "openpyxl is unavailable")
class XlsxAnonymizeTests(unittest.TestCase):
    def test_anonymize_writes_safe_xlsx_and_saves_values_to_database(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            input_path = Path(tmp) / "plan.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet["A1"] = "Ответственный"
            sheet["B1"] = "Петров А.В."
            sheet["A2"] = "Узел"
            sheet["B2"] = "ВРН-Север-17"
            sheet["A3"] = "IP"
            sheet["B3"] = "10.74.16.1/24"
            workbook.save(input_path)

            code, stdout, stderr, document_id, db_path = run_anonymize(input_path, OfficeFakeDetector())

            self.assertEqual(code, 0)
            safe_path = safe_path_for(input_path, document_id)
            self.assertTrue(safe_path.exists())

            safe_workbook = load_workbook(safe_path)
            safe_sheet = safe_workbook["Main"]
            values = [safe_sheet["B1"].value, safe_sheet["B2"].value, safe_sheet["B3"].value]
            self.assertEqual(values, ["⟦PERSON_001⟧", "⟦OBJECT_001⟧", "⟦IP_001⟧"])
            self.assertEqual(load_mapping(document_id, db_path).entities["OBJECT_001"], "ВРН-Север-17")
            self.assertEqual(stderr, "")


@unittest.skipUnless(PPTX_AVAILABLE, "python-pptx is unavailable")
class PptxAnonymizeTests(unittest.TestCase):
    def test_anonymize_writes_safe_pptx_and_saves_values_to_database(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            input_path = Path(tmp) / "deck.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = "Объект ВРН-Север-17"
            text_box = slide.shapes.add_textbox(914400, 1371600, 7315200, 914400)
            text_box.text = "Инженер Петров А.В. настроил 10.74.16.1/24"
            presentation.save(input_path)

            code, stdout, stderr, document_id, db_path = run_anonymize(input_path, OfficeFakeDetector())

            self.assertEqual(code, 0)
            safe_path = safe_path_for(input_path, document_id)
            self.assertTrue(safe_path.exists())

            safe_presentation = Presentation(safe_path)
            safe_text = "\n".join(_pptx_text(safe_presentation))
            self.assertNotIn("Петров А.В.", safe_text)
            self.assertNotIn("ВРН-Север-17", safe_text)
            self.assertNotIn("10.74.16.1/24", safe_text)
            self.assertIn("⟦PERSON_001⟧", safe_text)
            self.assertIn("⟦OBJECT_001⟧", safe_text)
            self.assertIn("⟦IP_001⟧", safe_text)
            self.assertEqual(load_mapping(document_id, db_path).entities["PERSON_001"], "Петров А.В.")
            self.assertEqual(stderr, "")


def _docx_text(document):
    for paragraph in document.paragraphs:
        yield paragraph.text
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield cell.text


def _pptx_text(presentation):
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                yield shape.text


if __name__ == "__main__":
    unittest.main()