import tempfile
import unittest
from pathlib import Path

from core.database import (
    DocumentNotFoundError,
    find_documents_by_placeholders,
    init_db,
    save_document,
)
from core.document_lookup import (
    AmbiguousDocumentError,
    placeholder_keys_in,
    resolve_document_id,
)
from core.models import Mapping


try:
    from docx import Document
except ImportError:  # pragma: no cover - exercised only without the optional test dependency
    Document = None


class DocumentLookupTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.db_path = self.tmp / "test.db"
        init_db(self.db_path)

    def tearDown(self):
        self._tmp.cleanup()

    def save(self, document_id, entities):
        save_document(
            document_id,
            Mapping.create(source_name=f"{document_id}.txt", entities=entities),
            self.db_path,
        )

    def test_placeholder_keys_are_extracted_from_text(self):
        keys = placeholder_keys_in("⟦PERSON_001⟧ и ⟦IP_002⟧, снова ⟦PERSON_001⟧, [[NOT_1]]")

        self.assertEqual(keys, {"PERSON_001", "IP_002"})

    def test_search_returns_documents_owning_all_found_placeholders(self):
        self.save("AAAA0001", {"PERSON_001": "a", "IP_001": "b"})
        self.save("BBBB0002", {"PERSON_001": "c"})

        matches = find_documents_by_placeholders({"PERSON_001", "IP_001"}, self.db_path)

        self.assertEqual([match.id for match in matches], ["AAAA0001"])
        self.assertTrue(matches[0].is_complete)

    def test_number_in_file_name_wins(self):
        self.save("AAAA0001", {"PERSON_001": "a"})
        self.save("BBBB0002", {"PERSON_001": "b"})

        document_id = resolve_document_id(
            self.tmp / "report_SAFE_bbbb0002_edited.txt",
            "⟦PERSON_001⟧",
            self.db_path,
        )

        self.assertEqual(document_id, "BBBB0002")

    def test_unknown_number_in_file_name_falls_back_to_placeholders(self):
        self.save("AAAA0001", {"PERSON_001": "a", "IP_001": "b"})

        document_id = resolve_document_id(
            self.tmp / "report_SAFE_FFFFFFFF.txt",
            "⟦PERSON_001⟧ ⟦IP_001⟧",
            self.db_path,
        )

        self.assertEqual(document_id, "AAAA0001")

    def test_complete_match_is_preferred_over_partial(self):
        self.save("AAAA0001", {"PERSON_001": "a"})
        self.save("BBBB0002", {"PERSON_001": "b", "IP_001": "c"})

        document_id = resolve_document_id(self.tmp / "answer.txt", "⟦PERSON_001⟧", self.db_path)

        self.assertEqual(document_id, "AAAA0001")

    def test_single_partial_match_is_used_so_restore_can_report_missing(self):
        self.save("AAAA0001", {"PERSON_001": "a", "IP_001": "b"})

        document_id = resolve_document_id(self.tmp / "answer.txt", "⟦PERSON_001⟧", self.db_path)

        self.assertEqual(document_id, "AAAA0001")

    def test_several_complete_matches_are_ambiguous(self):
        self.save("AAAA0001", {"PERSON_001": "a"})
        self.save("BBBB0002", {"PERSON_001": "b"})

        with self.assertRaises(AmbiguousDocumentError) as raised:
            resolve_document_id(self.tmp / "answer.txt", "⟦PERSON_001⟧", self.db_path)

        self.assertEqual({c.id for c in raised.exception.candidates}, {"AAAA0001", "BBBB0002"})

    def test_explicit_choice_resolves_ambiguity(self):
        self.save("AAAA0001", {"PERSON_001": "a"})
        self.save("BBBB0002", {"PERSON_001": "b"})

        document_id = resolve_document_id(
            self.tmp / "answer.txt", "⟦PERSON_001⟧", self.db_path, explicit_id="bbbb0002"
        )

        self.assertEqual(document_id, "BBBB0002")

    def test_file_without_placeholders_is_not_found(self):
        self.save("AAAA0001", {"PERSON_001": "a"})

        with self.assertRaises(DocumentNotFoundError):
            resolve_document_id(self.tmp / "answer.txt", "обычный текст", self.db_path)

    @unittest.skipUnless(Document is not None, "python-docx is unavailable")
    def test_number_in_docx_properties_is_used(self):
        from core.office import write_document_id

        self.save("AAAA0001", {"PERSON_001": "a"})
        self.save("BBBB0002", {"PERSON_001": "b"})
        path = self.tmp / "answer.docx"
        document = Document()
        document.add_paragraph("⟦PERSON_001⟧")
        document.save(path)
        write_document_id(path, "BBBB0002")

        document_id = resolve_document_id(path, "⟦PERSON_001⟧", self.db_path)

        self.assertEqual(document_id, "BBBB0002")


if __name__ == "__main__":
    unittest.main()