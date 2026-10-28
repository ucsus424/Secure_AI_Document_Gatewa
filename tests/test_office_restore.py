import io
import re
import tempfile
import unittest
from pathlib import Path

import app
from core.database import init_db, save_document
from core.models import Entity, Mapping


try:
    from docx import Document
except ImportError:  # pragma: no cover - exercised only without the optional test dependency
    Document = None


@unittest.skipUnless(Document is not None, "python-docx is unavailable")
class DocxRestoreTests(unittest.TestCase):
    def test_restore_writes_final_docx_and_restores_document_parts(self):
        with tempfile.TemporaryDirectory() as tmp:
            safe_path = Path(tmp) / "result_SAFE_AB12CD34.docx"
            db_path = Path(tmp) / "test.db"

            document = Document()
            document.add_paragraph("Инженер ⟦PERSON_001⟧ прибыл на ⟦OBJECT_001⟧.")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Ответственный: ⟦PERSON_001⟧"
            document.sections[0].header.paragraphs[0].text = "Объект ⟦OBJECT_001⟧"
            document.save(safe_path)
            _save_values(
                db_path,
                "AB12CD34",
                {
                    "PERSON_001": "Петров А.В.",
                    "OBJECT_001": "ВРН-Север-17",
                },
            )

            stdout = io.StringIO()
            stderr = io.StringIO()
            code = app.run(
                ["restore", "--db", str(db_path), str(safe_path), "AB12CD34"],
                stdout=stdout,
                stderr=stderr,
            )

            self.assertEqual(code, 0)
            final_path = Path(tmp) / "result_FINAL.docx"
            self.assertTrue(final_path.exists())
            final_document = Document(final_path)
            self.assertEqual(
                final_document.paragraphs[0].text,
                "Инженер Петров А.В. прибыл на ВРН-Север-17.",
            )
            self.assertEqual(final_document.tables[0].cell(0, 0).text, "Ответственный: Петров А.В.")
            self.assertEqual(
                final_document.sections[0].header.paragraphs[0].text,
                "Объект ВРН-Север-17",
            )
            self.assertNotIn("⟦", "\n".join(_docx_text(final_document)))
            self.assertIn(str(final_path), stdout.getvalue())
            self.assertEqual(stderr.getvalue(), "")

    def test_restore_handles_placeholder_split_across_runs(self):
        with tempfile.TemporaryDirectory() as tmp:
            safe_path = Path(tmp) / "split_SAFE_AB12CD34.docx"
            db_path = Path(tmp) / "test.db"

            document = Document()
            paragraph = document.add_paragraph("Инженер ")
            paragraph.add_run("⟦PERSON_")
            paragraph.add_run("001⟧")
            paragraph.add_run(" на объекте.")
            document.save(safe_path)
            _save_values(db_path, "AB12CD34", {"PERSON_001": "Петров А.В."})

            code = app.run(
                ["restore", "--db", str(db_path), str(safe_path), "AB12CD34"],
                stdout=io.StringIO(),
                stderr=io.StringIO(),
            )

            self.assertEqual(code, 0)
            final_document = Document(Path(tmp) / "split_FINAL.docx")
            self.assertEqual(final_document.paragraphs[0].text, "Инженер Петров А.В. на объекте.")

    def test_anonymize_restore_round_trip_preserves_run_formatting(self):
        class PersonDetector:
            def detect_text(self, text):
                return [Entity(type="PERSON", value="Петров А.В.", confidence=0.99)]

        with tempfile.TemporaryDirectory() as tmp:
            input_path = Path(tmp) / "styled.docx"
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run("Перед ")
            first_name_run = paragraph.add_run("Петров")
            first_name_run.bold = True
            initials_run = paragraph.add_run(" А.В.")
            initials_run.bold = True
            tail_run = paragraph.add_run(" после")
            tail_run.italic = True
            document.save(input_path)

            db_path = Path(tmp) / "test.db"
            anonymize_stdout = io.StringIO()
            anonymize_code = app.run(
                ["anonymize", "--db", str(db_path), str(input_path)],
                detector_factory=lambda args: PersonDetector(),
                stdout=anonymize_stdout,
                stderr=io.StringIO(),
            )
            self.assertEqual(anonymize_code, 0)
            document_id = re.search(r"Document ID: ([0-9A-F]{8})", anonymize_stdout.getvalue()).group(1)

            restore_code = app.run(
                [
                    "restore",
                    "--db",
                    str(db_path),
                    str(Path(tmp) / f"styled_SAFE_{document_id}.docx"),
                    document_id,
                ],
                stdout=io.StringIO(),
                stderr=io.StringIO(),
            )
            self.assertEqual(restore_code, 0)

            final_document = Document(Path(tmp) / "styled_FINAL.docx")
            final_paragraph = final_document.paragraphs[0]
            self.assertEqual(final_paragraph.text, "Перед Петров А.В. после")
            self.assertTrue(
                any(run.text == "Петров А.В." and run.bold for run in final_paragraph.runs)
            )
            self.assertTrue(
                any(run.text == " после" and run.italic for run in final_paragraph.runs)
            )

    def test_restore_missing_placeholder_fails_without_final_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            safe_path = Path(tmp) / "result_SAFE_AB12CD34.docx"
            db_path = Path(tmp) / "test.db"

            document = Document()
            document.add_paragraph("Инженер ⟦PERSON_001⟧ прибыл на объект.")
            document.save(safe_path)
            _save_values(
                db_path,
                "AB12CD34",
                {
                    "PERSON_001": "Петров А.В.",
                    "OBJECT_001": "ВРН-Север-17",
                },
            )

            stdout = io.StringIO()
            stderr = io.StringIO()
            code = app.run(
                ["restore", "--db", str(db_path), str(safe_path), "AB12CD34"],
                stdout=stdout,
                stderr=stderr,
            )

            self.assertNotEqual(code, 0)
            self.assertFalse((Path(tmp) / "result_FINAL.docx").exists())
            self.assertIn("OBJECT_001", stderr.getvalue())
            self.assertNotIn("Петров А.В.", stderr.getvalue())
            self.assertEqual(stdout.getvalue(), "")


def _save_values(db_path, document_id, entities):
    init_db(db_path)
    save_document(
        document_id,
        Mapping.create(source_name="input.docx", entities=entities),
        db_path,
    )


def _docx_text(document):
    for paragraph in document.paragraphs:
        yield paragraph.text
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield cell.text
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            yield paragraph.text


if __name__ == "__main__":
    unittest.main()