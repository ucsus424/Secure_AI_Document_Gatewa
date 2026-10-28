import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

from core.database import init_db, load_mapping, save_document
from core.models import Entity, Mapping
from web_ui import GatewayWebApp


try:
    from docx import Document
except ImportError:  # pragma: no cover - exercised only without the optional test dependency
    Document = None


DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class FakeDetector:
    def detect_text(self, text):
        return [Entity(type="PERSON", value="Петров А.В.", confidence=0.99)]


class EmptyDetector:
    def detect_text(self, text):
        return []


class RecordingDetector(FakeDetector):
    def __init__(self):
        self.calls = []

    def detect_text(self, text):
        self.calls.append(text)
        return super().detect_text(text)


class WebUiTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.db_path = Path(self._tmp.name) / "test.db"
        self.app = self.make_app(FakeDetector())

    def tearDown(self):
        self._tmp.cleanup()

    def make_app(self, detector):
        return GatewayWebApp(detector_factory=lambda: detector, db_path=self.db_path)

    def save_values(self, document_id, entities):
        init_db(self.db_path)
        save_document(
            document_id,
            Mapping.create(source_name="input.docx", entities=entities),
            self.db_path,
        )

    def count_documents(self):
        with sqlite3.connect(self.db_path) as connection:
            return connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0]

    def test_home_page_exposes_local_anonymize_and_restore_workflows(self):
        status, headers, body = _request(self.app, "GET", "/")

        html = body.decode("utf-8")
        self.assertEqual(status, "200 OK")
        self.assertIn("text/html", headers["Content-Type"])
        self.assertIn("Создать безопасную копию", html)
        self.assertIn("Восстановить DOCX", html)
        self.assertIn('name="document_id"', html)
        self.assertNotIn("mapping.json", html)
        self.assertIn("127.0.0.1", html)

    def test_anonymize_returns_only_safe_file_and_saves_values_to_database(self):
        status, headers, body = _request(
            self.app,
            "POST",
            "/anonymize",
            files={
                "document": ("input.txt", "Инженер Петров А.В. на объекте.".encode("utf-8"), "text/plain"),
            },
        )

        self.assertEqual(status, "200 OK")
        self.assertIn("text/plain", headers["Content-Type"])
        document_id = headers["X-Document-Id"]
        self.assertRegex(document_id, r"^[0-9A-F]{8}$")
        self.assertIn(f"input_SAFE_{document_id}.txt", headers["Content-Disposition"])
        self.assertEqual(body.decode("utf-8"), "Инженер ⟦PERSON_001⟧ на объекте.")

        mapping = load_mapping(document_id, self.db_path)
        self.assertEqual(mapping.entities["PERSON_001"], "Петров А.В.")

    def test_anonymize_automatically_chunks_large_upload(self):
        detector = RecordingDetector()
        app = self.make_app(detector)
        source = (
            ("обычная строка без чувствительных данных\n" * 400)
            + "Инженер Петров А.В. на объекте."
        ).encode("utf-8")

        status, headers, body = _request(
            app,
            "POST",
            "/anonymize",
            files={"document": ("large.txt", source, "text/plain")},
        )

        self.assertEqual(status, "200 OK")
        self.assertGreater(len(detector.calls), 1)
        self.assertTrue(all(len(chunk) <= 8000 for chunk in detector.calls))
        self.assertIn("⟦PERSON_001⟧", body.decode("utf-8"))

    def test_anonymize_rejects_empty_detection_result_and_saves_nothing(self):
        app = self.make_app(EmptyDetector())

        status, headers, body = _request(
            app,
            "POST",
            "/anonymize",
            files={
                "document": ("input.txt", "Контакт: user@example.com".encode("utf-8"), "text/plain"),
            },
        )

        payload = json.loads(body.decode("utf-8"))
        self.assertEqual(status, "502 Bad Gateway")
        self.assertIn("application/json", headers["Content-Type"])
        self.assertNotIn("user@example.com", payload["error"])
        self.assertEqual(self.count_documents(), 0)

    def test_restore_txt_by_document_id(self):
        self.save_values("AB12CD34", {"PERSON_001": "Петров А.В."})

        status, headers, body = _request(
            self.app,
            "POST",
            "/restore",
            files={"document": ("result_SAFE_AB12CD34.txt", "⟦PERSON_001⟧ на объекте.".encode("utf-8"), "text/plain")},
            fields={"document_id": "ab12cd34"},
        )

        self.assertEqual(status, "200 OK")
        self.assertIn("result_FINAL.txt", headers["Content-Disposition"])
        self.assertEqual(body.decode("utf-8"), "Петров А.В. на объекте.")

    def test_restore_requires_document_id(self):
        status, headers, body = _request(
            self.app,
            "POST",
            "/restore",
            files={"document": ("result_SAFE.txt", "⟦PERSON_001⟧".encode("utf-8"), "text/plain")},
        )

        self.assertEqual(status, "400 Bad Request")
        self.assertIn("номер", json.loads(body.decode("utf-8"))["error"])

    def test_restore_unknown_document_id_returns_not_found(self):
        status, headers, body = _request(
            self.app,
            "POST",
            "/restore",
            files={"document": ("result_SAFE.txt", "⟦PERSON_001⟧".encode("utf-8"), "text/plain")},
            fields={"document_id": "FFFFFFFF"},
        )

        self.assertEqual(status, "404 Not Found")
        self.assertIn("не найден", json.loads(body.decode("utf-8"))["error"])

    def test_anonymize_then_restore_round_trip(self):
        original = "Инженер Петров А.В. на объекте."
        _, headers, safe_body = _request(
            self.app,
            "POST",
            "/anonymize",
            files={"document": ("input.txt", original.encode("utf-8"), "text/plain")},
        )

        status, headers, body = _request(
            self.app,
            "POST",
            "/restore",
            files={"document": ("answer.txt", safe_body, "text/plain")},
            fields={"document_id": headers["X-Document-Id"]},
        )

        self.assertEqual(status, "200 OK")
        self.assertEqual(body.decode("utf-8"), original)

    @unittest.skipUnless(Document is not None, "python-docx is unavailable")
    def test_restore_returns_final_docx(self):
        self.save_values("AB12CD34", {"PERSON_001": "Петров А.В."})
        safe_bytes = io.BytesIO()
        document = Document()
        document.add_paragraph("Инженер ⟦PERSON_001⟧ на объекте.")
        document.save(safe_bytes)

        status, headers, body = _request(
            self.app,
            "POST",
            "/restore",
            files={"document": ("result_SAFE_AB12CD34.docx", safe_bytes.getvalue(), DOCX_TYPE)},
            fields={"document_id": "AB12CD34"},
        )

        self.assertEqual(status, "200 OK")
        self.assertEqual(headers["Content-Type"], DOCX_TYPE)
        self.assertIn("result_FINAL.docx", headers["Content-Disposition"])
        restored = Document(io.BytesIO(body))
        self.assertEqual(restored.paragraphs[0].text, "Инженер Петров А.В. на объекте.")

    @unittest.skipUnless(Document is not None, "python-docx is unavailable")
    def test_restore_missing_placeholder_returns_safe_error(self):
        self.save_values(
            "AB12CD34",
            {"PERSON_001": "Петров А.В.", "OBJECT_001": "ВРН-Север-17"},
        )
        safe_bytes = io.BytesIO()
        document = Document()
        document.add_paragraph("Инженер ⟦PERSON_001⟧ на объекте.")
        document.save(safe_bytes)

        status, headers, body = _request(
            self.app,
            "POST",
            "/restore",
            files={"document": ("result_SAFE_AB12CD34.docx", safe_bytes.getvalue(), DOCX_TYPE)},
            fields={"document_id": "AB12CD34"},
        )

        payload = json.loads(body.decode("utf-8"))
        self.assertEqual(status, "422 Unprocessable Entity")
        self.assertIn("application/json", headers["Content-Type"])
        self.assertIn("OBJECT_001", payload["error"])
        self.assertNotIn("Петров А.В.", payload["error"])


def _request(app, method, path, files=None, fields=None):
    body = b""
    content_type = ""
    if files or fields:
        body, content_type = _multipart_body(files or {}, fields or {})

    environ = {
        "REQUEST_METHOD": method,
        "PATH_INFO": path,
        "CONTENT_LENGTH": str(len(body)),
        "CONTENT_TYPE": content_type,
        "wsgi.input": io.BytesIO(body),
        "wsgi.url_scheme": "http",
        "SERVER_NAME": "127.0.0.1",
        "SERVER_PORT": "8765",
    }
    captured = {}

    def start_response(status, headers):
        captured["status"] = status
        captured["headers"] = dict(headers)

    response_body = b"".join(app(environ, start_response))
    return captured["status"], captured["headers"], response_body


def _multipart_body(files, fields):
    boundary = "----SecureGatewayTestBoundary"
    chunks = []
    for field_name, value in fields.items():
        chunks.extend(
            [
                f"--{boundary}\r\n".encode("ascii"),
                f'Content-Disposition: form-data; name="{field_name}"\r\n\r\n'.encode("utf-8"),
                value.encode("utf-8"),
                b"\r\n",
            ]
        )
    for field_name, (filename, data, content_type) in files.items():
        chunks.extend(
            [
                f"--{boundary}\r\n".encode("ascii"),
                (
                    f'Content-Disposition: form-data; name="{field_name}"; '
                    f'filename="{filename}"\r\n'
                ).encode("utf-8"),
                f"Content-Type: {content_type}\r\n\r\n".encode("ascii"),
                data,
                b"\r\n",
            ]
        )
    chunks.append(f"--{boundary}--\r\n".encode("ascii"))
    return b"".join(chunks), f"multipart/form-data; boundary={boundary}"


if __name__ == "__main__":
    unittest.main()